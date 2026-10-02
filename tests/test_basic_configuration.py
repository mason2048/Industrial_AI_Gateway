import io
import json
import time

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from backend.database import Database
from backend.history import History
from backend.main import create_app
from backend.models import Tag
from backend.tag_manager import export_excel, import_excel


def test_individual_history_intervals_and_disabled_point(tmp_path):
    db = Database(tmp_path / "history.db")
    tags = [Tag(id=1, address="A", name="fast", type="FLOAT", threshold=100,
                history_interval_seconds=60),
            Tag(id=2, address="B", name="default", type="FLOAT", threshold=100),
            Tag(id=3, address="C", name="live only", type="FLOAT", save=False)]
    db.replace_tags(tags)
    history = History(db, heartbeat=1800)
    start = time.time()
    readings = {tag.id: {"value": 1, "quality": "Good"} for tag in tags}
    assert history.save(tags, readings, "simulation", now=start, monotonic=100) == 2
    assert history.save(tags, readings, "simulation", now=start + 59, monotonic=159) == 0
    assert history.save(tags, readings, "simulation", now=start + 60, monotonic=160) == 1
    assert history.save(tags, readings, "simulation", now=start + 1800, monotonic=1900) == 2
    with db.connect() as conn:
        assert conn.execute("SELECT COUNT(*) FROM history_data WHERE tag_id=3").fetchone()[0] == 0


def test_five_decimal_threshold_accumulates_from_saved_baseline(tmp_path):
    db = Database(tmp_path / "history.db")
    tag = Tag(id=1, address="A", name="vacuum", type="FLOAT", threshold=.00005)
    db.replace_tags([tag])
    history = History(db)
    now = time.time()
    def save(value, delta):
        return history.save([tag], {1: {"value": value, "quality": "Good"}},
                            "simulation", now=now + delta)
    assert save(1.0, 0) == 1
    for index in range(1, 5):
        assert save(1 + index / 100000, index) == 0
    assert save(1.00005, 5) == 1


def test_excel_new_columns_and_legacy_defaults():
    tag = Tag(id=1, address="A", name="vacuum", type="FLOAT", threshold=.00001,
              history_interval_seconds=900, precision=5)
    assert import_excel(export_excel([tag])) == [tag]
    workbook = load_workbook(io.BytesIO(export_excel([tag])))
    workbook.active.delete_cols(12, 3)
    stream = io.BytesIO()
    workbook.save(stream)
    legacy = import_excel(stream.getvalue())[0]
    assert legacy.history_interval_seconds is None and legacy.precision == 5


def test_periodic_only_policy_does_not_save_every_scan(tmp_path):
    db = Database(tmp_path / "history.db")
    tag = Tag(id=1, address="A", name="periodic", type="FLOAT", record_changes=False,
              history_interval_seconds=1800, threshold=0)
    db.replace_tags([tag])
    history = History(db)
    now = time.time()
    def save(value, delta, quality="Good"):
        return history.save([tag], {1: {"value": value, "quality": quality}},
                            "simulation", now=now + delta)
    assert save(1, 0) == 1
    assert save(2, 1) == 0
    assert save(3, 1799) == 0
    assert save(4, 1800) == 1
    assert save(None, 1801, "BadCommunicationError") == 1


def test_runtime_settings_are_authorized_versioned_and_persisted(tmp_path):
    (tmp_path / "config").mkdir()
    (tmp_path / "frontend").mkdir()
    (tmp_path / "config/config.json").write_text(json.dumps({"mode": "simulation",
        "endpoint": "opc.tcp://127.0.0.1:4840", "backup_enabled": False}))
    app = create_app(tmp_path)
    with TestClient(app) as client:
        config = client.get("/api/config")
        body = {"mode": "simulation", "endpoint": config.json()["endpoint"],
                "poll_interval": .5, "batch_size": 50, "heartbeat_seconds": 900,
                "retention_days": 7}
        assert client.post("/api/connection", json=body).status_code == 403
        headers = {"X-Operator-Pin": (tmp_path / "data/operator_pin.txt").read_text(),
                   "If-Match": config.headers["ETag"]}
        updated = client.post("/api/connection", json=body, headers=headers)
        assert updated.status_code == 200, updated.text
        assert app.state.gateway.config["poll_interval"] == .5
        assert app.state.history.heartbeat == 900 and app.state.history.retention == 7
        stored = json.loads((tmp_path / "config/config.json").read_text())
        assert stored["batch_size"] == 50 and stored["heartbeat_seconds"] == 900
        assert client.post("/api/connection", json=body, headers=headers).status_code == 412
        assert client.post("/api/connection", json={**body, "retention_days": 8},
                           headers=headers).status_code == 422
        template = client.get("/api/tags/template")
        assert template.status_code == 200
        wb = load_workbook(io.BytesIO(template.content))
        assert wb.active.max_row == 1
        assert "保存间隔秒" in [cell.value for cell in wb.active[1]]


@pytest.mark.parametrize("change", [{"history_interval_seconds": 0}, {"precision": 11},
                                     {"history_interval_seconds": float("inf")}])
def test_invalid_per_point_settings(change):
    with pytest.raises(ValueError):
        Tag(id=1, address="A", name="vacuum", type="FLOAT", **change)
