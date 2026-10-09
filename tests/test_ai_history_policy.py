"""Per-point AI history preferences filter saved evidence only, on demand."""
from datetime import datetime, timedelta, timezone
import io
import json
import time

from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook
import pytest

from backend.ai_api import LocalDataProvider, router
from backend.database import Database
from backend.llm import AIConfigStore, ModelService
from backend.main import create_app
from backend.models import AIQuery, Tag
from backend.observability import Redactor
from backend.tag_manager import export_excel, import_excel
from tests.test_llm_integration import fake_model

BASE = datetime(2026, 10, 9, tzinfo=timezone.utc)


def point(tag_id=1, **changes):
    return {"id": tag_id, "revision": 1, "name": f"压力{tag_id}", "device": "Pump", "address": f"A{tag_id}",
            "type": "FLOAT", "unit": "Pa", "save": True, "threshold": .5, "precision": 5,
            "ai_description": "原始定义描述", "ai_history_mode": "changes", "ai_history_interval_seconds": None,
            "history_interval_seconds": None, "value": 1, "quality": "Good", "timestamp": BASE.isoformat(), **changes}


def rows(values, seconds, *, tag_id=1, revision=1, qualities=None, kind="FLOAT"):
    return [{"id": tag_id * 1000 + index, "tag_id": tag_id, "tag_revision": revision,
             "name": f"压力{tag_id}", "device": "Pump", "source": "opcua", "connection_id": "plant-a",
             "data_type": kind, "unit": "Pa", "value": value,
             "quality": qualities[index] if qualities else "Good", "timestamp": (BASE + timedelta(seconds=second)).isoformat()}
            for index, (value, second) in enumerate(zip(values, seconds))]


def read(all_rows):
    def query(**filters):
        selected = list(all_rows)
        for key in ("device", "variable", "source", "connection_id", "tag_revision"):
            if filters.get(key) is not None:
                field = "name" if key == "variable" else key
                selected = [row for row in selected if row[field] == filters[key]]
        selected.sort(key=lambda item: (item["timestamp"], item["id"]))
        limit, offset = filters.get("limit", 2000), filters.get("offset", 0)
        return {"items": selected[offset:offset + limit], "total": len(selected), "limit": limit, "offset": offset,
                "has_more": offset + limit < len(selected), "start": BASE.isoformat(),
                "end": (BASE + timedelta(days=1)).isoformat(), "summary": [], "note": "已保存样本"}
    return query


def provider(points, samples, *, definitions=None, default=1800):
    snapshot = lambda device=None: {"items": [p for p in points if not device or p["device"] == device],
                                   "mode": "opcua", "connection_id": "plant-a", "connected": True,
                                   "good": len(points), "total": len(points), "storage_error": ""}
    definition = lambda tag_id, revision: (definitions or {}).get((tag_id, revision))
    return LocalDataProvider(snapshot, read(samples), definition, default_interval=lambda: default)


def times(result):
    return [int((datetime.fromisoformat(row["timestamp"]) - BASE).total_seconds()) for row in result["items"]]


def test_changes_and_interval_produce_different_saved_evidence():
    samples = rows([1, 1, 2, 2, 2, 3], [0, 10, 20, 30, 60, 90])
    changes = provider([point()], samples).history_data(changed_only=True, use_point_settings=True)
    interval = provider([point(ai_history_mode="interval", ai_history_interval_seconds=30)], samples).history_data(
        changed_only=True, use_point_settings=True)
    assert times(changes) == [0, 20, 90]
    assert times(interval) == [0, 30, 60, 90]
    assert [row["change_reason"] for row in interval["items"]] == ["baseline", "interval_sample", "interval_sample", "interval_sample"]
    assert all(row["ai_history_mode"] == "interval" and row["ai_history_interval_seconds"] == 30 for row in interval["items"])
    assert interval["filter"]["point_settings"] and interval["filter"]["point_policies"][0]["interval_source"] == "ai_history_interval"


@pytest.mark.parametrize("ai_interval,saved_interval,global_interval,expected,origin", [
    (20, 30, 90, 20, "ai_history_interval"), (None, 30, 90, 30, "history_interval"),
    (None, None, 90, 90, "global_default")])
def test_interval_inherits_point_then_global(ai_interval, saved_interval, global_interval, expected, origin):
    selected = point(ai_history_mode="interval", ai_history_interval_seconds=ai_interval,
                     history_interval_seconds=saved_interval)
    result = provider([selected], rows([1] * 10, list(range(0, 100, 10))), default=global_interval).history_data(
        changed_only=True, use_point_settings=True)
    assert times(result) == list(range(0, 100, expected))
    policy = result["filter"]["point_policies"][0]
    assert policy["interval_seconds"] == expected and policy["interval_source"] == origin


def test_latest_point_preference_applies_to_old_revision_without_rewriting_threshold_or_identity():
    old = point(threshold=1)
    current = point(revision=2, threshold=100, ai_description="当前描述", ai_history_mode="interval", ai_history_interval_seconds=30)
    samples = rows([0, .5, 1, 1], [0, 10, 30, 60])
    selected = provider([current], samples, definitions={(1, 1): old}).history_data(changed_only=True, use_point_settings=True)
    assert times(selected) == [0, 30, 60]
    assert all(row["tag_revision"] == 1 and row["threshold"] == 1 and row["ai_description"] == "原始定义描述" for row in selected["items"])
    assert all(row["ai_history_policy_revision"] == 2 and row["ai_history_policy_source"] == "current_point" for row in selected["items"])
    current["ai_history_mode"] = "changes"
    changes = provider([current], samples, definitions={(1, 1): old}).history_data(changed_only=True, use_point_settings=True)
    assert times(changes) == [0, 30]  # Uses threshold=1 from the old definition.


def test_retired_point_falls_back_to_historical_definition_preference():
    retired = point(ai_history_mode="interval", ai_history_interval_seconds=45)
    result = provider([], rows([1] * 5, [0, 30, 50, 90, 100]), definitions={(1, 1): retired}).history_data(
        changed_only=True, use_point_settings=True)
    assert times(result) == [0, 50, 100]
    assert result["filter"]["point_policies"][0]["policy_source"] == "historical_definition"


def test_mixed_point_policies_preserve_quality_changes_without_fake_periodic_records():
    samples = rows([1, 1, 2, 2], [0, 10, 20, 40]) + rows([False, True, None, True, True, False],
        [0, 10, 12, 13, 40, 43], tag_id=2, qualities=["Good", "Good", "BadCommunicationError", "Good", "Good", "Good"], kind="BOOL")
    result = provider([point(), point(2, type="BOOL", ai_history_mode="interval", ai_history_interval_seconds=30)], samples).history_data(
        changed_only=True, use_point_settings=True)
    first = {**result, "items": [row for row in result["items"] if row["tag_id"] == 1]}
    second = {**result, "items": [row for row in result["items"] if row["tag_id"] == 2]}
    assert times(first) == [0, 20]
    assert times(second) == [0, 12, 13, 43]
    assert [row["change_reason"] for row in second["items"]] == ["baseline", "quality_change", "quality_change", "interval_sample"]
    assert {row["ai_history_mode"] for row in result["items"]} == {"changes", "interval"}
    assert len(result["filter"]["point_policies"]) == 2


def test_time_mode_picks_next_existing_sample_and_paginate_after_filter():
    data = provider([point(ai_history_mode="interval", ai_history_interval_seconds=30)], rows([1] * 3, [0, 20, 100]))
    result = data.history_data(changed_only=True, use_point_settings=True, limit=1, offset=1)
    assert times(result) == [100] and result["total"] == 2 and not result["has_more"]
    assert "不补造采样" in result["note"]


def test_legacy_change_filter_and_raw_opt_out_remain_compatible():
    samples = rows([1] * 4, [0, 10, 30, 60])
    data = provider([point(ai_history_mode="interval", ai_history_interval_seconds=30)], samples)
    legacy = data.history_data(changed_only=True)
    assert times(legacy) == [0] and not legacy["filter"]["point_settings"]
    assert legacy["filter"]["point_policies"][0]["policy_source"] == "legacy_changes"
    raw = data.history_data(changed_only=False, use_point_settings=True)
    assert times(raw) == [0, 10, 30, 60] and "filter" not in raw


def test_get_history_defaults_to_point_settings_with_explicit_legacy_switch():
    data = provider([point(ai_history_mode="interval", ai_history_interval_seconds=30)], rows([1] * 4, [0, 10, 30, 60]))
    app = FastAPI()
    app.include_router(router(data))
    with TestClient(app) as client:
        selected = client.get("/api/ai/history").json()
        legacy = client.get("/api/ai/history?use_point_settings=false").json()
        raw = client.get("/api/ai/history?changed_only=false").json()
    assert times(selected) == [0, 30, 60] and selected["filter"]["point_settings"]
    assert times(legacy) == [0] and times(raw) == [0, 10, 30, 60]


def test_interval_scan_cap_and_invalid_timestamps_are_disclosed():
    data = provider([point(ai_history_mode="interval", ai_history_interval_seconds=1)], rows([1] * 5, [0, 1, 2, 3, 4]))
    capped = data.history_data(changed_only=True, use_point_settings=True, scan_limit=3)
    assert capped["truncated"] and not capped["filter"]["complete"] and capped["filter"]["scanned"] == 3
    samples = rows([1] * 3, [0, 30, 60])
    samples[0]["timestamp"] = "invalid"
    # Preserve rows when time cannot be compared; never invent a missing time.
    result = provider([point(ai_history_mode="interval", ai_history_interval_seconds=30)], samples).history_data(
        changed_only=True, use_point_settings=True)
    assert result["filter"]["invalid_timestamps"] == [{"tag_id": 1, "tag_revision": 1}]
    assert any(row["change_reason"] == "timestamp_unavailable" for row in result["items"])
    assert "历史采集时间无效" in result["note"]


def test_unsaved_point_uses_current_even_with_interval_preference(tmp_path):
    data = provider([point(save=False, ai_history_mode="interval", ai_history_interval_seconds=30)], [])
    def must_not_read(**filters):
        raise AssertionError("Live-only question cannot read SQLite history")
    data.history = must_not_read
    service = ModelService(AIConfigStore(tmp_path), data, Redactor([]))
    result = data.analyze(AIQuery(question="压力1多少"))
    assert result["query_type"] == "current"
    context, _, _ = service._context(AIQuery(question="压力1多少"))
    assert context["query_type"] == "current" and context["history_items"] == []


@pytest.mark.parametrize("changes", [{"ai_history_mode": "upload"}, {"ai_history_interval_seconds": 0},
                                    {"ai_history_interval_seconds": -1}, {"ai_history_interval_seconds": 604801},
                                    {"ai_history_interval_seconds": float("inf")}, {"ai_history_interval_seconds": float("nan")}])
def test_invalid_policy_values_rejected(changes):
    with pytest.raises(ValueError):
        Tag(id=1, address="A", name="压力", type="FLOAT", **changes)


def test_excel_roundtrip_labels_validation_and_old_template_defaults():
    tag = Tag(id=1, address="A", name="压力", type="FLOAT", ai_history_mode="interval", ai_history_interval_seconds=900)
    content = export_excel([tag])
    assert import_excel(content) == [tag]
    workbook = load_workbook(io.BytesIO(content))
    sheet = workbook.active
    assert sheet.cell(1, 15).value == "AI取数方式" and sheet.cell(1, 16).value == "AI取数间隔秒"
    assert sheet.cell(2, 15).value == "时间" and sheet.cell(2, 16).value == 900
    assert any('"变化,时间"' == validation.formula1 for validation in sheet.data_validations.dataValidation)
    sheet.cell(2, 15).value = "自动上传"
    stream = io.BytesIO()
    workbook.save(stream)
    with pytest.raises(ValueError, match="AI取数方式必须为变化或时间"):
        import_excel(stream.getvalue())
    sheet.delete_cols(15, 2)
    stream = io.BytesIO()
    workbook.save(stream)
    old = import_excel(stream.getvalue())[0]
    assert old.ai_history_mode == "changes" and old.ai_history_interval_seconds is None


def installation(tmp_path):
    for folder in ("config", "frontend", "data"):
        (tmp_path / folder).mkdir()
    (tmp_path / "config/config.json").write_text(json.dumps({"mode": "simulation", "poll_interval": .1,
        "heartbeat_seconds": 90, "backup_enabled": False}))
    (tmp_path / "frontend/index.html").write_text("<title>policy test</title>")


def test_old_database_defaults_without_silent_definition_migration(tmp_path):
    installation(tmp_path)
    db = Database(tmp_path / "data/history.db")
    legacy = Tag(id=1, address="A", name="压力", type="FLOAT").model_dump(exclude={"ai_history_mode", "ai_history_interval_seconds"})
    db.replace_tags([legacy])
    before = db.tags()
    app = create_app(tmp_path)
    with TestClient(app) as client:
        actual = client.get("/api/current").json()["items"][0]
        assert actual["ai_history_mode"] == "changes" and actual["ai_history_interval_seconds"] is None
        assert app.state.db.tags() == before
    assert Database(tmp_path / "data/history.db").tags() == before


def test_saved_preference_restart_llm_payload_and_no_background_model_requests(tmp_path):
    installation(tmp_path)
    db = Database(tmp_path / "data/history.db")
    tag = Tag(id=1, address="A", name="压力", device="Pump", type="FLOAT", threshold=100,
              record_changes=False, history_interval_seconds=60)
    db.replace_tags([tag])
    app = create_app(tmp_path)
    with fake_model() as (url, calls):
        with TestClient(app) as client:
            pin = {"X-Operator-Pin": app.state.operator_pin}
            model = client.get("/api/ai/config")
            assert client.post("/api/ai/config", headers={**pin, "If-Match": model.headers["ETag"]},
                json={"provider": "ollama", "model": "mock-model", "base_url": url}).status_code == 200
            connection = app.state.config_store.snapshot()["connection_id"]
            base = time.time() - 200
            db.insert_history([(1, base + second, 1, "Good", "simulation", "Pump", "压力", "Pa", "FLOAT",
                                connection, 1, None, None, None) for second in (0, 15, 30, 45, 60, 90, 120)])
            fetched = client.get("/api/tags")
            edited = fetched.json()
            edited[0].update(ai_history_mode="interval", ai_history_interval_seconds=45)
            assert client.put("/api/tags", headers={**pin, "If-Match": fetched.headers["ETag"]}, json=edited).status_code == 200
            assert calls == []  # Neither save nor acquisition triggers model work.
            model = client.get("/api/ai/config")
            question = {"question": "历史压力数据", "start": datetime.fromtimestamp(base, timezone.utc).isoformat(),
                        "end": datetime.fromtimestamp(base + 130, timezone.utc).isoformat()}
            result = client.post("/api/ai/query", headers={**pin, "If-Match": model.headers["ETag"]}, json=question)
            assert result.status_code == 200, result.text
            evidence = result.json()["evidence"]
            selected = evidence["history_items"]
            assert len(selected) == 3 and all(row["ai_history_mode"] == "interval" and row["tag_revision"] == 1 for row in selected)
            assert evidence["history_filter"]["point_settings"]
            sent = json.loads(calls[0]["body"]["messages"][1]["content"])["read_only_evidence"]
            assert sent == evidence
            assert "变化阈值或时间间隔" in calls[0]["body"]["messages"][0]["content"]
            assert len(calls) == 1
        restored = create_app(tmp_path)
        with TestClient(restored) as client:
            stored = client.get("/api/tags").json()[0]
            assert stored["ai_history_mode"] == "interval" and stored["ai_history_interval_seconds"] == 45
            assert restored.state.model_service.store.public()["provider"] == "ollama"
            assert len(calls) == 1  # Restart remains idle.
            default = client.get("/api/ai/history", params={"start": question["start"], "end": question["end"]}).json()
            assert len(default["items"]) == 3 and default["filter"]["point_policies"][0]["interval_seconds"] == 45
