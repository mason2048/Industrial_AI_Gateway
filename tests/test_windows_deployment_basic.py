"""Deploy/packaging tests use temporary files; no field installation is modified."""
import hashlib
import json
from pathlib import Path
import shutil
import sys
from types import SimpleNamespace
import uuid
import zipfile

import pytest

from scripts import bootstrap
from scripts.package_release import ROOT, TOP_LEVEL, build_release


def prepared_environment(tmp_path, monkeypatch):
    interpreter = tmp_path / ".venv" / ("Scripts/python.exe" if sys.platform == "win32" else "bin/python")
    interpreter.parent.mkdir(parents=True)
    interpreter.touch()
    (tmp_path / "requirements-lock.txt").write_text("example==1\n")
    monkeypatch.setattr(bootstrap.subprocess, "check_output", lambda *a, **k: "3.12\n")
    return interpreter


def test_offline_install_disables_index_and_uses_quoted_argument_paths(tmp_path, monkeypatch):
    interpreter = prepared_environment(tmp_path, monkeypatch)
    wheels = tmp_path / "offline wheels"
    wheels.mkdir()
    calls = []
    monkeypatch.setattr(bootstrap.subprocess, "run", lambda command, **kwargs: calls.append(command))
    assert bootstrap.ensure_environment(tmp_path, offline=True, wheelhouse=wheels) == interpreter
    assert calls[0][:5] == [str(interpreter), "-m", "pip", "install", "--no-index"]
    assert calls[0][5:7] == ["--find-links", str(wheels)]
    assert calls[1] == [str(interpreter), "-m", "pip", "check"]
    # A successful offline install starts again without reaching a package server.
    bootstrap.ensure_environment(tmp_path)
    assert len(calls) == 2


def test_missing_offline_wheels_never_records_success(tmp_path, monkeypatch):
    prepared_environment(tmp_path, monkeypatch)
    calls = []
    monkeypatch.setattr(bootstrap.subprocess, "run", lambda command, **kwargs: calls.append(command))
    with pytest.raises(RuntimeError, match="Wheel directory"):
        bootstrap.ensure_environment(tmp_path, offline=True)
    assert calls == []
    assert not (tmp_path / ".venv/dependencies.sha256").exists()


def test_service_bundle_also_marks_base_dependencies_ready(tmp_path, monkeypatch):
    interpreter = tmp_path / ".venv/Scripts/python.exe"
    interpreter.parent.mkdir(parents=True)
    interpreter.touch()
    (tmp_path / "requirements-lock.txt").write_text("example==1\n")
    (tmp_path / "requirements-windows-lock.txt").write_text("-r requirements-lock.txt\npywin32==311\n")
    monkeypatch.setattr(bootstrap, "os", SimpleNamespace(name="nt"))
    monkeypatch.setattr(bootstrap.subprocess, "check_output", lambda *a, **k: "3.12\n")
    calls = []
    monkeypatch.setattr(bootstrap.subprocess, "run", lambda command, **kwargs: calls.append(command))
    bootstrap.ensure_environment(tmp_path, windows_service=True)
    assert (tmp_path / ".venv/dependencies-windows.sha256").is_file()
    assert (tmp_path / ".venv/dependencies.sha256").is_file()
    bootstrap.ensure_environment(tmp_path)
    assert len(calls) == 2  # Starting the foreground entry after service setup remains offline.


def test_unknown_install_option_is_rejected_before_environment_changes(monkeypatch):
    calls = []
    monkeypatch.setattr(bootstrap, "ensure_environment", lambda **kwargs: calls.append(kwargs))
    assert bootstrap.main(["init", "--ofline"]) == 1
    assert calls == []


@pytest.fixture
def package_source(tmp_path):
    source = tmp_path / "source"
    source.mkdir()
    for name in TOP_LEVEL:
        shutil.copy2(ROOT / name, source / name)
    for name in ("backend", "scripts", "frontend", "docs", "tests"):
        shutil.copytree(ROOT / name, source / name, ignore=shutil.ignore_patterns("__pycache__"))
    for path in ROOT.glob("*.bat"):
        shutil.copy2(path, source / path.name)
    secret = "field-private-" + uuid.uuid4().hex
    (source / "data/backups").mkdir(parents=True)
    (source / "data/history.db").write_bytes(secret.encode())
    (source / "data/operator_pin.txt").write_text(secret)
    (source / "data/backups/private.txt").write_text(secret)
    (source / "config").mkdir()
    (source / "config/config.json").write_text(json.dumps({"username": secret}))
    (source / "config/client-key.pem").write_text(secret)
    return source


def test_release_excludes_field_data_and_contains_verified_empty_template(package_source, tmp_path):
    from openpyxl import load_workbook
    import io
    output = tmp_path / "release.zip"
    secret = (package_source / "data/operator_pin.txt").read_bytes()
    result = build_release(output, package_source)
    assert result["sha256"] == hashlib.sha256(output.read_bytes()).hexdigest()
    assert Path(result["checksum"]).read_text().startswith(result["sha256"])
    with zipfile.ZipFile(output) as archive:
        prefix = "Industrial_AI_Gateway/"
        names = [name.removeprefix(prefix) for name in archive.namelist()]
        assert "data/" in names
        assert not any(name.startswith("data/") and name != "data/" for name in names)
        assert not any(".venv" in name or "__pycache__" in name for name in names)
        assert not any(secret in archive.read(name) for name in archive.namelist())
        config = json.loads(archive.read(prefix + "config/config.json"))
        assert config["mode"] == "opcua" and config["username"] == ""
        assert not config["simulation_write_enabled"]
        manifest = json.loads(archive.read(prefix + "RELEASE_MANIFEST.json"))
        assert not manifest["contains_field_data"]
        for name, details in manifest["files"].items():
            content = archive.read(prefix + name)
            assert details == {"size": len(content), "sha256": hashlib.sha256(content).hexdigest()}
        workbook = load_workbook(io.BytesIO(archive.read(prefix + "templates/PLC点位导入模板.xlsx")))
        try:
            assert workbook.active.max_row == 1
            assert "NodeId" in next(workbook.active.values)
        finally:
            workbook.close()


def test_release_rejects_linked_source_file(package_source, tmp_path):
    linked = package_source / "scripts/leak.py"
    try:
        linked.symlink_to(package_source / "config/client-key.pem")
    except OSError:
        pytest.skip("Creating a symlink requires local permission on this Windows account")
    with pytest.raises(ValueError, match="symbolic link"):
        build_release(tmp_path / "unsafe.zip", package_source)
    assert not (tmp_path / "unsafe.zip").exists()


def test_service_status_reports_state_name_without_starting_gateway(monkeypatch, capsys):
    import backend.windows_service as service
    states = {"SERVICE_STOPPED": 1, "SERVICE_START_PENDING": 2, "SERVICE_STOP_PENDING": 3,
              "SERVICE_RUNNING": 4, "SERVICE_CONTINUE_PENDING": 5, "SERVICE_PAUSE_PENDING": 6,
              "SERVICE_PAUSED": 7}
    monkeypatch.setattr(service, "os", SimpleNamespace(name="nt"))
    monkeypatch.setitem(sys.modules, "servicemanager", SimpleNamespace())
    monkeypatch.setitem(sys.modules, "win32service", SimpleNamespace(**states))
    monkeypatch.setitem(sys.modules, "win32serviceutil", SimpleNamespace(
        ServiceFramework=object, QueryServiceStatus=lambda name: (0, states["SERVICE_RUNNING"])))
    assert service.main(["status"]) == 0
    assert "IndustrialAIGateway: RUNNING" in capsys.readouterr().out
    with pytest.raises(SystemExit, match="Usage"):
        service.main(["start", "unexpected"])
