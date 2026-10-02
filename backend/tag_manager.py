import io
import zipfile
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from .models import Tag

HEADERS = {"ID":"id", "地址":"address", "名称":"name", "类型":"type", "单位":"unit", "权限":"permission", "AI描述":"ai_description", "保存":"save", "阈值":"threshold", "设备":"device", "NodeId":"node_id", "保存间隔秒":"history_interval_seconds", "小数位数":"precision", "记录变化":"record_changes"}


def validate_tags(items):
    tags = [x if isinstance(x, Tag) else Tag.model_validate(x) for x in items]
    if not 1 <= len(tags) <= 1000:
        raise ValueError("点位数量必须为1至1000")
    for key in (lambda t:t.id, lambda t:(t.device,t.name), lambda t:(t.device,t.address)):
        values = [key(t) for t in tags]
        if len(values) != len(set(values)):
            raise ValueError("ID、同设备名称和同设备地址不能重复")
    nodes = [(t.device,t.node_id) for t in tags if t.node_id]
    if len(nodes) != len(set(nodes)):
        raise ValueError("同设备NodeId不能重复")
    return tags


def import_excel(content):
    if len(content) > 5 * 1024 * 1024:
        raise ValueError("Excel文件不能超过5MB")
    with zipfile.ZipFile(io.BytesIO(content)) as z:
        if sum(x.file_size for x in z.infolist()) > 25 * 1024 * 1024:
            raise ValueError("Excel解压内容过大")
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
    try:
        ws = wb.active
        if ws.max_row > 1001 or ws.max_column > 20:
            raise ValueError("最多1000行点位、20列；请移除多余格式空行")
        rows = ws.iter_rows(values_only=True)
        header = [str(x).strip() if x is not None else "" for x in next(rows)]
        if len([x for x in header if x]) != len(set(x for x in header if x)):
            raise ValueError("表头重复")
        if not {"ID","地址","名称","类型","单位","权限","AI描述","保存","阈值"} <= set(header):
            raise ValueError("缺少必需中文表头，请使用导出模板")
        tags = []
        for number, row in enumerate(rows, 2):
            if all(x is None for x in row):
                continue
            try:
                d = {HEADERS[h]:v for h,v in zip(header,row) if h in HEADERS and v is not None}
                flag = str(d.get("save", "YES")).strip().upper()
                if flag not in ("YES","NO","TRUE","FALSE","1","0"):
                    raise ValueError("保存字段必须YES或NO")
                d["save"] = flag in ("YES","TRUE","1")
                if d.get("threshold") in (None,"","-"):
                    d["threshold"] = 0
                if d.get("history_interval_seconds") in (None,"","-"):
                    d["history_interval_seconds"] = None
                if d.get("precision") in (None,"","-"):
                    d["precision"] = 5
                if "记录变化" in header:
                    flag = str(d.get("record_changes", "NO")).strip().upper()
                    if flag not in ("YES", "NO", "TRUE", "FALSE", "1", "0"):
                        raise ValueError("记录变化字段必须YES或NO")
                    d["record_changes"] = flag in ("YES", "TRUE", "1")
                tags.append(Tag.model_validate(d))
            except Exception as exc:
                raise ValueError(f"Excel第{number}行：{exc}") from exc
        return validate_tags(tags)
    finally:
        wb.close()


def export_excel(tags):
    wb = Workbook()
    ws = wb.active
    ws.title = "PLC点位"
    ws.append(list(HEADERS))
    for tag in tags:
        d = (tag if isinstance(tag, Tag) else Tag.model_validate(tag)).model_dump()
        row = [d[k] if k not in ("save", "record_changes") else ("YES" if d[k] else "NO") for k in HEADERS.values()]
        ws.append(row)
    for row in ws:
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = "s"  # Treat user text as text, never Excel formulas.
            cell.alignment = Alignment(vertical="center")
        ws.row_dimensions[row[0].row].height = 24
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor="133D45")
        cell.font = Font(color="FFFFFF", bold=True)
    for col, width in zip("ABCDEFGHIJKLMN", [10,24,22,12,12,12,38,12,18,22,42,18,14,16]):
        ws.column_dimensions[col].width = width
    for row in range(2, ws.max_row+1):
        if row % 2 == 0:
            for cell in ws[row]: cell.fill = PatternFill("solid", fgColor="EDF5F4")
        ws.cell(row,9).number_format = "0.00000E+00"
    for col, options in [("D","BOOL,WORD,DWORD,FLOAT"),("F","READ,WRITE"),("H","YES,NO"),("N","YES,NO")]:
        v = DataValidation(type="list", formula1=f'"{options}"')
        v.errorTitle = "请选择有效值"
        v.showErrorMessage = True
        ws.add_data_validation(v)
        v.add(f"{col}2:{col}1001")
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = ws.dimensions
    stream = io.BytesIO()
    wb.save(stream)
    return stream.getvalue()
